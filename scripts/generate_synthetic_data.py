# ABOUTME: Generates synthetic PBM pharmacy claim files for Databricks ACQ ingest testing.
# ABOUTME: Reads column mappings from Excel, infers field types, writes pipe-delimited files.

import sys
import os
import random
import string
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(
    0, os.path.expanduser("~/AppData/Roaming/Python/Python314/site-packages")
)
import openpyxl

SEED = 42
ROWS_PER_FILE = 1000
NULL_RATE = 0.08
DELIMITER = "|"

PROJECT_ROOT = Path(__file__).resolve().parent.parent
EXCEL_PATH = (
    PROJECT_ROOT
    / "docs"
    / "project"
    / "DLH PharmacyClaimTransaction to ACQ table column mappings.xlsx"
)
OUTPUT_DIR = PROJECT_ROOT / "data" / "synthetic"


# ---------------------------------------------------------------------------
# Shared entity pools
# ---------------------------------------------------------------------------

FIRST_NAMES = [
    "James", "Mary", "Robert", "Patricia", "John", "Jennifer", "Michael",
    "Linda", "David", "Elizabeth", "William", "Barbara", "Richard", "Susan",
    "Joseph", "Jessica", "Thomas", "Sarah", "Christopher", "Karen", "Charles",
    "Lisa", "Daniel", "Nancy", "Matthew", "Betty", "Anthony", "Margaret",
    "Mark", "Sandra", "Donald", "Ashley", "Steven", "Dorothy", "Paul", "Kimberly",
    "Andrew", "Emily", "Joshua", "Donna", "Kenneth", "Michelle", "Kevin", "Carol",
    "Brian", "Amanda", "George", "Melissa", "Timothy", "Deborah",
]

LAST_NAMES = [
    "Smith", "Johnson", "Williams", "Brown", "Jones", "Garcia", "Miller",
    "Davis", "Rodriguez", "Martinez", "Hernandez", "Lopez", "Gonzalez",
    "Wilson", "Anderson", "Thomas", "Taylor", "Moore", "Jackson", "Martin",
    "Lee", "Perez", "Thompson", "White", "Harris", "Sanchez", "Clark",
    "Ramirez", "Lewis", "Robinson", "Walker", "Young", "Allen", "King",
    "Wright", "Scott", "Torres", "Nguyen", "Hill", "Flores", "Green",
    "Adams", "Nelson", "Baker", "Hall", "Rivera", "Campbell", "Mitchell",
    "Carter", "Roberts",
]

DRUG_NAMES = [
    "ATORVASTATIN", "LISINOPRIL", "METFORMIN HCL", "AMLODIPINE BESYLATE",
    "METOPROLOL SUCCINATE", "OMEPRAZOLE", "SIMVASTATIN", "LOSARTAN POTASSIUM",
    "ALBUTEROL SULFATE", "GABAPENTIN", "HYDROCHLOROTHIAZIDE", "SERTRALINE HCL",
    "MONTELUKAST SODIUM", "ESCITALOPRAM OXALATE", "ROSUVASTATIN CALCIUM",
    "LEVOTHYROXINE SODIUM", "PANTOPRAZOLE SODIUM", "FLUTICASONE PROPIONATE",
    "TRAMADOL HCL", "PREDNISONE", "AMOXICILLIN", "TRAZODONE HCL",
    "TAMSULOSIN HCL", "DULOXETINE HCL", "CLOPIDOGREL BISULFATE",
    "MELOXICAM", "CARVEDILOL", "VENLAFAXINE HCL", "BUPROPION HCL",
    "CYCLOBENZAPRINE HCL", "INSULIN GLARGINE", "ADALIMUMAB", "APIXABAN",
    "ELIQUIS", "HUMIRA", "OZEMPIC", "JARDIANCE", "TRULICITY",
    "ENTRESTO", "XARELTO",
]

PHARMACY_NAMES = [
    "CVS PHARMACY", "WALGREENS", "WALMART PHARMACY", "RITE AID",
    "KROGER PHARMACY", "COSTCO PHARMACY", "PUBLIX PHARMACY",
    "ALBERTSONS PHARMACY", "HEB PHARMACY", "MEIJER PHARMACY",
    "SAFEWAY PHARMACY", "SAM'S CLUB PHARMACY", "TARGET PHARMACY",
    "WEGMANS PHARMACY", "GIANT EAGLE PHARMACY", "STOP & SHOP PHARMACY",
    "RITE AID PHARMACY", "KINNEY DRUGS", "HARMONS PHARMACY",
    "GOOD NEIGHBOR PHARMACY", "HEALTH MART", "MEDICINE SHOPPE",
    "CARE PHARMACY", "HOMETOWN PHARMACY", "GENOA HEALTHCARE",
    "OPTUMRX MAIL", "CVS SPECIALTY", "ACCREDO SPECIALTY",
    "BRIOVA SPECIALTY", "DIPLOMAT SPECIALTY",
]

STATES = [
    "AL", "AK", "AZ", "AR", "CA", "CO", "CT", "DE", "FL", "GA",
    "HI", "ID", "IL", "IN", "IA", "KS", "KY", "LA", "ME", "MD",
    "MA", "MI", "MN", "MS", "MO", "MT", "NE", "NV", "NH", "NJ",
    "NM", "NY", "NC", "ND", "OH", "OK", "OR", "PA", "RI", "SC",
    "SD", "TN", "TX", "UT", "VT", "VA", "WA", "WV", "WI", "WY",
]

CITIES = [
    "NEW YORK", "LOS ANGELES", "CHICAGO", "HOUSTON", "PHOENIX",
    "PHILADELPHIA", "SAN ANTONIO", "SAN DIEGO", "DALLAS", "JACKSONVILLE",
    "AUSTIN", "COLUMBUS", "CHARLOTTE", "INDIANAPOLIS", "SAN FRANCISCO",
    "SEATTLE", "DENVER", "NASHVILLE", "OKLAHOMA CITY", "PORTLAND",
    "MILWAUKEE", "ALBUQUERQUE", "TUCSON", "FRESNO", "MESA",
    "ATLANTA", "OMAHA", "MIAMI", "TULSA", "TAMPA",
]

STREETS = [
    "MAIN ST", "OAK AVE", "MAPLE DR", "ELM ST", "CEDAR LN",
    "PINE RD", "WALNUT ST", "BIRCH CT", "SPRUCE WAY", "ASH BLVD",
    "PARK AVE", "LAKE DR", "HILL RD", "VALLEY LN", "RIVER ST",
    "FOREST DR", "MEADOW LN", "SUNRISE BLVD", "SUNSET DR", "HIGHLAND AVE",
]

THERAPEUTIC_CLASSES = [
    "CARDIOVASCULAR", "ANTIDIABETIC", "CNS", "RESPIRATORY",
    "GASTROINTESTINAL", "ANTI-INFECTIVE", "MUSCULOSKELETAL",
    "ENDOCRINE", "IMMUNOLOGIC", "HEMATOLOGIC",
]

GPI_PREFIXES = [
    "39", "27", "72", "44", "49", "12", "66", "23", "65", "83",
]

DAW_CODES = ["0", "1", "2", "3", "4", "5", "7", "8", "9"]

CLAIM_STATUSES = ["P", "R", "A"]
CLAIM_STATUS_WEIGHTS = [0.85, 0.10, 0.05]

RELATIONSHIP_CODES = ["01", "02", "03", "04"]
GENDER_CODES = ["M", "F", "U"]

PBM_SHORT_NAMES = {
    "CAP NAVITUS": "cap_navitus",
    "BMICH": "bmich",
    "CVS_4000": "cvs_4000",
    "EMPIRX": "empirx",
    "EMPIRX_REJECT": "empirx_reject",
    "ESI_2300": "esi_2300",
    "ESI_HMRK": "esi_hmrk",
    "NAVITUS": "navitus",
    "BCKC": "bckc",
    "OPTUM_CHF7": "optum_chf7",
    "PRIME MAGELLAN": "prime_magellan",
    "RIGHTWAY": "rightway",
}


def build_entity_pools(rng):
    members = []
    for i in range(50):
        fn = FIRST_NAMES[i % len(FIRST_NAMES)]
        ln = LAST_NAMES[i % len(LAST_NAMES)]
        dob = date(1940, 1, 1) + timedelta(days=rng.randint(0, 23725))
        gender = rng.choice(["M", "F"])
        state = rng.choice(STATES)
        city = rng.choice(CITIES)
        zip5 = f"{rng.randint(10000, 99999)}"
        phone = f"{rng.randint(200, 999)}{rng.randint(2000000, 9999999)}"
        street_num = rng.randint(1, 9999)
        street = rng.choice(STREETS)
        mbi = "".join(
            rng.choices(string.ascii_uppercase + string.digits, k=11)
        )
        pbm_ids = {}
        for pbm in PBM_SHORT_NAMES.values():
            pbm_ids[pbm] = f"{rng.randint(100000000, 999999999)}"
        subscriber_id = f"SUB{rng.randint(10000000, 99999999)}"
        members.append({
            "first_name": fn,
            "last_name": ln,
            "middle_initial": fn[0],
            "dob": dob,
            "gender": gender,
            "state": state,
            "city": city,
            "zip": zip5,
            "phone": phone,
            "address1": f"{street_num} {street}",
            "address2": "",
            "mbi": mbi,
            "pbm_ids": pbm_ids,
            "subscriber_id": subscriber_id,
            "subscriber_first": fn,
            "subscriber_last": ln,
            "relationship": "01",
        })

    ndcs = []
    for i in range(200):
        labeler = f"{rng.randint(0, 99999):05d}"
        product = f"{rng.randint(0, 9999):04d}"
        package = f"{rng.randint(0, 99):02d}"
        ndc11 = f"{labeler}{product}{package}"
        drug_name = DRUG_NAMES[i % len(DRUG_NAMES)]
        gpi = f"{rng.choice(GPI_PREFIXES)}{rng.randint(100000, 999999):06d}"
        tc = rng.choice(THERAPEUTIC_CLASSES)
        awp = round(rng.uniform(0.10, 500.00), 4)
        wac = round(awp * rng.uniform(0.60, 0.95), 4)
        generic_ind = rng.choice(["Y", "N"])
        ndcs.append({
            "ndc11": ndc11,
            "ndc_dash": f"{labeler}-{product}-{package}",
            "drug_name": drug_name,
            "gpi": gpi,
            "therapeutic_class": tc,
            "awp": awp,
            "wac": wac,
            "generic_indicator": generic_ind,
            "brand_generic": "GENERIC" if generic_ind == "Y" else "BRAND",
        })

    pharmacies = []
    for i in range(30):
        ncpdp = f"{rng.randint(1000000, 9999999)}"
        npi = f"1{rng.randint(100000000, 999999999)}"
        name = PHARMACY_NAMES[i % len(PHARMACY_NAMES)]
        state = rng.choice(STATES)
        city = rng.choice(CITIES)
        zip5 = f"{rng.randint(10000, 99999)}"
        phone = f"{rng.randint(200, 999)}{rng.randint(2000000, 9999999)}"
        chain_ind = rng.choice(["C", "I"])
        specialty = rng.choice(["Y", "N"])
        pharmacies.append({
            "ncpdp": ncpdp,
            "npi": npi,
            "name": name,
            "state": state,
            "city": city,
            "zip": zip5,
            "phone": phone,
            "address": f"{rng.randint(1, 9999)} {rng.choice(STREETS)}",
            "chain_independent": chain_ind,
            "specialty": specialty,
        })

    prescribers = []
    for i in range(20):
        npi = f"1{rng.randint(100000000, 999999999)}"
        fn = FIRST_NAMES[(i + 25) % len(FIRST_NAMES)]
        ln = LAST_NAMES[(i + 25) % len(LAST_NAMES)]
        dea = f"A{ln[0].upper()}{rng.randint(1000000, 9999999)}"
        state = rng.choice(STATES)
        city = rng.choice(CITIES)
        zip5 = f"{rng.randint(10000, 99999)}"
        specialty_code = f"{rng.randint(100, 999)}"
        prescribers.append({
            "npi": npi,
            "first_name": fn,
            "last_name": ln,
            "middle_initial": fn[0],
            "dea": dea,
            "state": state,
            "city": city,
            "zip": zip5,
            "address": f"{rng.randint(1, 9999)} {rng.choice(STREETS)}",
            "phone": f"{rng.randint(200, 999)}{rng.randint(2000000, 9999999)}",
            "specialty_code": specialty_code,
            "gender": rng.choice(["M", "F"]),
        })

    hierarchies = []
    for i in range(5):
        company = f"CO{i + 1:03d}"
        division = f"DIV{i + 1:03d}"
        client = f"CL{i + 1:04d}"
        cdc_id = f"{company}-{division}-{client}"
        hierarchies.append({
            "company_code": company,
            "division_code": division,
            "client_code": client,
            "cdc_id": cdc_id,
        })

    benefit_groups = []
    for i in range(10):
        bg_id = f"BG{i + 1:06d}"
        bg_name = f"BENEFIT GROUP {i + 1}"
        lob = rng.choice(["COMMERCIAL", "MEDICARE", "MEDICAID", "EXCHANGE"])
        benefit_groups.append({
            "id": bg_id,
            "name": bg_name,
            "lob": lob,
        })

    return members, ndcs, pharmacies, prescribers, hierarchies, benefit_groups


# ---------------------------------------------------------------------------
# Field-type inference and value generation
# ---------------------------------------------------------------------------

def classify_field(name):
    upper = name.upper()

    if upper in (
        "COMPANY_DIVISION_CLIENT_ID", "COMPANY_CODE", "DIVISION_CODE",
        "CLIENT_CODE",
    ):
        return "hierarchy"
    if upper == "PHARMACY_CLAIM_TRANSACTION_INTERNAL_ID":
        return "claim_internal_id"
    if upper == "PHARMACY_CLAIM_TRANSACTION_ACTIVE_INDICATOR":
        return "active_indicator"

    if "SUBSCRIBER" in upper and ("FIRST" in upper or "FNAME" in upper):
        return "subscriber_first_name"
    if "SUBSCRIBER" in upper and ("LAST" in upper or "LNAME" in upper):
        return "subscriber_last_name"
    if "SUBSCRIBER" in upper and "MIDDLE" in upper:
        return "subscriber_middle_initial"
    if "SUBSCRIBER" in upper and "PHONE" in upper:
        return "subscriber_phone"
    if "SUBSCRIBER" in upper and "ADDRESS" in upper and "2" in upper:
        return "address2"
    if "SUBSCRIBER" in upper and "ADDRESS" in upper:
        return "subscriber_address"
    if "SUBSCRIBER" in upper and ("CITY" in upper or "CTY" in upper):
        return "subscriber_city"
    if "SUBSCRIBER" in upper and "STATE" in upper:
        return "subscriber_state"
    if "SUBSCRIBER" in upper and "ZIP" in upper:
        return "subscriber_zip"
    if upper in ("SUBSCRIBER_ID", "SUBSCRIBER_NUMBER", "SUBSCRIBER_NUM",
                  "SUBSCRIBER_NBR", "SBSCRBR_ID"):
        return "subscriber_id"

    if "CARDHOLDER" in upper and ("FIRST" in upper or "FNAME" in upper):
        return "member_first_name"
    if "CARDHOLDER" in upper and ("LAST" in upper or "LNAME" in upper):
        return "member_last_name"
    if "CARDHOLDER" in upper and "MIDDLE" in upper:
        return "member_middle_initial"
    if "CARDHOLDER" in upper and ("DOB" in upper or "BIRTH" in upper):
        return "member_dob"
    if "CARDHOLDER" in upper and "GENDER" in upper:
        return "member_gender"
    if "CARDHOLDER_ID" in upper:
        return "member_pbm_id"

    if ("PATIENT" in upper or "MEMBER" in upper or "MEM_" in upper) and (
        "FIRST" in upper or "FNAME" in upper
    ):
        return "member_first_name"
    if ("PATIENT" in upper or "MEMBER" in upper or "MEM_" in upper) and (
        "LAST" in upper or "LNAME" in upper
    ):
        return "member_last_name"
    if ("PATIENT" in upper or "MEMBER" in upper) and "MIDDLE" in upper:
        return "member_middle_initial"
    if ("PATIENT" in upper or "MEMBER" in upper) and "SUFFIX" in upper:
        return "member_suffix"
    if upper in ("MEM_FIRST_NAME",):
        return "member_first_name"
    if upper in ("MEM_LAST_NAME",):
        return "member_last_name"
    if upper in ("MEM_ID", "MEMBER_ID", "MEMBER_NUMBER", "PATIENT_ID",
                  "PERSON_CODE", "PERSON_NUMBER"):
        return "member_pbm_id"
    if ("PATIENT" in upper or "MEMBER" in upper or "MEM_" in upper) and (
        "DOB" in upper or "BIRTH" in upper
    ):
        return "member_dob"
    if ("PATIENT" in upper or "MEMBER" in upper or "MEM_" in upper) and (
        "GENDER" in upper or "SEX" in upper
    ):
        return "member_gender"
    if ("PATIENT" in upper or "MEMBER" in upper) and "AGE" in upper:
        return "member_age"
    if ("PATIENT" in upper or "MEMBER" in upper) and "PHONE" in upper:
        return "member_phone"
    if ("PATIENT" in upper or "MEMBER" in upper) and "ADDRESS" in upper and "2" in upper:
        return "address2"
    if ("PATIENT" in upper or "MEMBER" in upper) and "ADDRESS" in upper:
        return "member_address"
    if ("PATIENT" in upper or "MEMBER" in upper) and (
        "CITY" in upper or "CTY" in upper
    ):
        return "member_city"
    if ("PATIENT" in upper or "MEMBER" in upper) and "STATE" in upper:
        return "member_state"
    if ("PATIENT" in upper or "MEMBER" in upper) and "ZIP" in upper:
        return "member_zip"
    if "RELATIONSHIP" in upper:
        return "relationship"

    if "PRESCRIBER" in upper or "PROVIDER" in upper:
        if "NPI" in upper:
            return "prescriber_npi"
        if "FIRST" in upper or "FNAME" in upper:
            return "prescriber_first_name"
        if "LAST" in upper or "LNAME" in upper:
            return "prescriber_last_name"
        if "MIDDLE" in upper:
            return "prescriber_middle_initial"
        if "DEA" in upper:
            return "prescriber_dea"
        if "ADDRESS" in upper and "2" in upper:
            return "address2"
        if "ADDRESS" in upper:
            return "prescriber_address"
        if "CITY" in upper or "CTY" in upper:
            return "prescriber_city"
        if "STATE" in upper:
            return "prescriber_state"
        if "ZIP" in upper:
            return "prescriber_zip"
        if "PHONE" in upper:
            return "prescriber_phone"
        if "DOB" in upper or "BIRTH" in upper:
            return "prescriber_dob"
        if "GENDER" in upper or "SEX" in upper:
            return "prescriber_gender"
        if "SPECIALTY" in upper or "TAXONOMY" in upper:
            return "prescriber_specialty"
        if "ID" in upper or "NUMBER" in upper or "NUM" in upper:
            return "prescriber_npi"

    if "PHARMACY" in upper:
        if "NPI" in upper:
            return "pharmacy_npi"
        if "NCPDP" in upper:
            return "pharmacy_ncpdp"
        if "NAME" in upper:
            return "pharmacy_name"
        if "ADDRESS" in upper and "2" in upper:
            return "address2"
        if "ADDRESS" in upper:
            return "pharmacy_address"
        if "CITY" in upper or "CTY" in upper:
            return "pharmacy_city"
        if "STATE" in upper:
            return "pharmacy_state"
        if "ZIP" in upper:
            return "pharmacy_zip"
        if "PHONE" in upper:
            return "pharmacy_phone"
        if "CHAIN" in upper:
            return "pharmacy_chain"
        if "SPECIALTY" in upper:
            return "pharmacy_specialty"
        if "ID" in upper or "NUMBER" in upper or "NUM" in upper:
            return "pharmacy_ncpdp"

    if "SERVICE_PROVIDER_NPI" in upper or upper == "PHARMACY_NPI":
        return "pharmacy_npi"
    if upper in ("NCPDP_ID", "SERVICE_PROVIDER_ID", "NABP_NUMBER",
                  "NABP_NUM", "NABP"):
        return "pharmacy_ncpdp"

    if upper in ("NDC", "NDC_NUMBER", "NDC_NUM", "NDC_11",
                  "NATIONAL_DRUG_CODE", "DRUG_NDC", "PRODUCT_ID_NDC"):
        return "ndc"
    if "DRUG_NAME" in upper or "DRUG_PRODUCT_NAME" in upper or "PRODUCT_NAME" in upper:
        return "drug_name"
    if "DRUG_GPI" in upper or upper == "GPI":
        return "drug_gpi"
    if "THERAPEUTIC" in upper:
        return "therapeutic_class"
    if "GENERIC" in upper and "INDICATOR" in upper:
        return "generic_indicator"
    if "BRAND_GENERIC" in upper or "DRUG_BRAND" in upper:
        return "brand_generic"
    if upper in ("AWP", "AWP_PRICE", "AVERAGE_WHOLESALE_UNIT_PRICE",
                  "AVERAGE_WHOLESALE_UNIT_PRICE_FULL", "AWP_UNIT_PRICE"):
        return "awp"
    if upper in ("WAC", "WAC_PRICE", "WAC_UNIT_PRICE"):
        return "wac"
    if "DAW" in upper or "DISPENSE_AS_WRITTEN" in upper:
        return "daw"

    if upper in ("CLAIM_NUMBER", "CLM_NBR", "CLM_NUM", "UNIQUE_CLAIM_NUMBER",
                  "RX_CLAIM_NUMBER", "CLAIM_ID", "CLAIM_REFERENCE_NUMBER",
                  "UNIQUE_CLAIM_REFERENCE_NUMBER"):
        return "claim_number"
    if upper in ("ORIGINAL_CLAIM_NUMBER", "PRIOR_AUTHORIZATION_NUMBER"):
        return "original_claim_number"
    if upper in ("CLAIM_STATUS", "CLAIM_STATUS_CODE", "CLM_STATUS",
                  "CLAIM_STATUS_TYPE"):
        return "claim_status"
    if upper in ("CLAIM_TYPE", "CLAIM_TYPE_CODE"):
        return "claim_type"
    if upper in ("RX_NUMBER", "PRESCRIPTION_NUMBER", "RX_NUM", "RX_NBR"):
        return "rx_number"
    if "REFILL" in upper:
        return "refill_number"
    if "LINE_NUMBER" in upper or "LINE_NUM" in upper:
        return "line_number"
    if "DAYS_SUPPLY" in upper or "DAYS_SUPPLIED" in upper:
        return "days_supply"
    if "QUANTITY" in upper or "QTY" in upper:
        return "quantity"

    if "GROUP" in upper and ("ID" in upper or "NUMBER" in upper or "NUM" in upper):
        return "group_id"
    if "GROUP" in upper and "NAME" in upper:
        return "group_name"
    if "BENEFIT" in upper and "GROUP" in upper:
        return "benefit_group"
    if "SEGMENT" in upper:
        return "segment_code"
    if "CONTRACT" in upper and "NUMBER" in upper:
        return "contract_number"
    if "CARRIER" in upper and ("ID" in upper or "NUMBER" in upper):
        return "carrier_id"
    if "CUSTOMER_ID" in upper:
        return "customer_id"
    if "SOURCE_SYSTEM" in upper:
        return "source_system"

    if upper in ("MBI", "MEDICARE_BENEFICIARY_IDENTIFIER",
                  "MEDICARE_BENEFICIARY_ID", "MEDICARE_ID"):
        return "mbi"
    if "MEDICARE" in upper:
        return "code"

    if "DATE" in upper or upper.endswith("_DT"):
        return "date_field"
    if upper.endswith("_DOB") or "BIRTH" in upper:
        return "member_dob"

    money_keywords = (
        "AMOUNT", "AMT", "COST", "PRICE", "FEE", "PAID", "COPAY",
        "BILLED", "CHARGE", "REIMBURS", "DEDUCTIBLE", "COINSURANCE",
        "OOP", "TAX", "DISCOUNT", "DISPENSING", "INGREDIENT",
        "INCENTIVE", "PENALTY", "WITHHOLD",
    )
    for kw in money_keywords:
        if kw in upper:
            return "money"

    if upper.endswith("_INDICATOR") or upper.endswith("_FLAG") or upper.endswith("_IND"):
        return "indicator"
    if upper.endswith("_GENDER") or upper.endswith("_SEX"):
        return "member_gender"
    if upper.endswith("_ZIP") or upper.endswith("_ZIP_CODE"):
        return "zip"
    if upper.endswith("_PHONE") or "TELEPHONE" in upper:
        return "phone"
    if upper.endswith("_STATE") or upper.endswith("_STATE_CODE"):
        return "state"
    if upper.endswith("_CITY") or upper.endswith("_CTY"):
        return "city"
    if "ADDRESS" in upper and "2" in upper:
        return "address2"
    if "ADDRESS" in upper:
        return "address"
    if upper.endswith("_NPI"):
        return "prescriber_npi"

    if upper.endswith("_CODE") or upper.endswith("_CD") or upper.endswith("_QUAL"):
        return "code"
    if upper.endswith("_TYPE") or upper.endswith("_QUALIFIER"):
        return "code"
    if upper.endswith("_ID") or upper.endswith("_NUMBER") or upper.endswith("_NUM"):
        return "numeric_id"
    if upper.endswith("_NBR"):
        return "numeric_id"
    if upper.endswith("_NAME") or "NAME" in upper:
        return "name_field"
    if upper.endswith("_STATUS"):
        return "code"
    if upper.endswith("_TEXT") or upper.endswith("_DESC") or upper.endswith("_DESCRIPTION"):
        return "text"
    if "TIMESTAMP" in upper or "INSERT" in upper or "UPDATE" in upper:
        return "timestamp"
    if "COUNT" in upper or "CNT" in upper:
        return "count"
    if "PERCENT" in upper or "PCT" in upper or "RATE" in upper:
        return "percent"
    if "RECORD" in upper and "INDICATOR" not in upper:
        return "code"

    return "generic"


class ValueGenerator:
    def __init__(self, rng, members, ndcs, pharmacies, prescribers,
                 hierarchies, benefit_groups, pbm_key):
        self.rng = rng
        self.members = members
        self.ndcs = ndcs
        self.pharmacies = pharmacies
        self.prescribers = prescribers
        self.hierarchies = hierarchies
        self.benefit_groups = benefit_groups
        self.pbm_key = pbm_key
        self._claim_counter = 0
        self._row_member = None
        self._row_ndc = None
        self._row_pharmacy = None
        self._row_prescriber = None
        self._row_hierarchy = None
        self._row_bg = None
        self._row_claim_num = None
        self._row_status = None

    def new_row(self):
        self._row_member = self.rng.choice(self.members)
        self._row_ndc = self.rng.choice(self.ndcs)
        self._row_pharmacy = self.rng.choice(self.pharmacies)
        self._row_prescriber = self.rng.choice(self.prescribers)
        self._row_hierarchy = self.rng.choice(self.hierarchies)
        self._row_bg = self.rng.choice(self.benefit_groups)
        self._claim_counter += 1
        self._row_claim_num = f"CLM{self._claim_counter:012d}"
        self._row_status = self.rng.choices(
            CLAIM_STATUSES, weights=CLAIM_STATUS_WEIGHTS, k=1
        )[0]

    def maybe_null(self, value, rate=NULL_RATE):
        if self.rng.random() < rate:
            return ""
        return value

    def generate(self, field_name, field_type):
        rng = self.rng
        m = self._row_member
        ndc = self._row_ndc
        pharm = self._row_pharmacy
        pres = self._row_prescriber
        hier = self._row_hierarchy
        bg = self._row_bg

        if field_type == "hierarchy":
            upper = field_name.upper()
            if "DIVISION_CLIENT" in upper:
                return hier["cdc_id"]
            if "COMPANY" in upper:
                return hier["company_code"]
            if "DIVISION" in upper:
                return hier["division_code"]
            if "CLIENT" in upper:
                return hier["client_code"]
            return hier["cdc_id"]

        if field_type == "claim_internal_id":
            return f"INT{self._claim_counter:012d}"

        if field_type == "active_indicator":
            return "Y"

        if field_type == "member_first_name":
            return m["first_name"]
        if field_type == "member_last_name":
            return m["last_name"]
        if field_type == "member_middle_initial":
            return m["middle_initial"]
        if field_type == "member_suffix":
            return self.maybe_null("")
        if field_type == "member_pbm_id":
            return m["pbm_ids"].get(self.pbm_key, m["pbm_ids"].get("navitus"))
        if field_type == "member_dob":
            return m["dob"].strftime("%Y-%m-%d")
        if field_type == "member_gender":
            return m["gender"]
        if field_type == "member_age":
            return str((date(2025, 6, 15) - m["dob"]).days // 365)
        if field_type == "member_phone":
            return m["phone"]
        if field_type == "member_address":
            return m["address1"]
        if field_type == "member_city":
            return m["city"]
        if field_type == "member_state":
            return m["state"]
        if field_type == "member_zip":
            return m["zip"]
        if field_type == "mbi":
            return self.maybe_null(m["mbi"], 0.30)

        if field_type == "subscriber_id":
            return m["subscriber_id"]
        if field_type == "subscriber_first_name":
            return m["subscriber_first"]
        if field_type == "subscriber_last_name":
            return m["subscriber_last"]
        if field_type == "subscriber_middle_initial":
            return m["middle_initial"]
        if field_type == "subscriber_phone":
            return m["phone"]
        if field_type == "subscriber_address":
            return m["address1"]
        if field_type == "subscriber_city":
            return m["city"]
        if field_type == "subscriber_state":
            return m["state"]
        if field_type == "subscriber_zip":
            return m["zip"]
        if field_type == "relationship":
            return rng.choice(RELATIONSHIP_CODES)

        if field_type == "prescriber_npi":
            return pres["npi"]
        if field_type == "prescriber_first_name":
            return pres["first_name"]
        if field_type == "prescriber_last_name":
            return pres["last_name"]
        if field_type == "prescriber_middle_initial":
            return pres["middle_initial"]
        if field_type == "prescriber_dea":
            return pres["dea"]
        if field_type == "prescriber_address":
            return pres["address"]
        if field_type == "prescriber_city":
            return pres["city"]
        if field_type == "prescriber_state":
            return pres["state"]
        if field_type == "prescriber_zip":
            return pres["zip"]
        if field_type == "prescriber_phone":
            return pres["phone"]
        if field_type == "prescriber_gender":
            return pres["gender"]
        if field_type == "prescriber_specialty":
            return pres["specialty_code"]
        if field_type == "prescriber_dob":
            dob = date(1950, 1, 1) + timedelta(days=rng.randint(0, 14600))
            return dob.strftime("%Y-%m-%d")

        if field_type == "pharmacy_npi":
            return pharm["npi"]
        if field_type == "pharmacy_ncpdp":
            return pharm["ncpdp"]
        if field_type == "pharmacy_name":
            return pharm["name"]
        if field_type == "pharmacy_address":
            return pharm["address"]
        if field_type == "pharmacy_city":
            return pharm["city"]
        if field_type == "pharmacy_state":
            return pharm["state"]
        if field_type == "pharmacy_zip":
            return pharm["zip"]
        if field_type == "pharmacy_phone":
            return pharm["phone"]
        if field_type == "pharmacy_chain":
            return pharm["chain_independent"]
        if field_type == "pharmacy_specialty":
            return pharm["specialty"]

        if field_type == "ndc":
            return ndc["ndc11"]
        if field_type == "drug_name":
            return ndc["drug_name"]
        if field_type == "drug_gpi":
            return ndc["gpi"]
        if field_type == "therapeutic_class":
            return ndc["therapeutic_class"]
        if field_type == "generic_indicator":
            return ndc["generic_indicator"]
        if field_type == "brand_generic":
            return ndc["brand_generic"]
        if field_type == "awp":
            return f"{ndc['awp']:.4f}"
        if field_type == "wac":
            return f"{ndc['wac']:.4f}"
        if field_type == "daw":
            return rng.choice(DAW_CODES)

        if field_type == "claim_number":
            return self._row_claim_num
        if field_type == "original_claim_number":
            if self._row_status == "R":
                return f"CLM{rng.randint(1, max(1, self._claim_counter - 1)):012d}"
            return ""
        if field_type == "claim_status":
            return self._row_status
        if field_type == "claim_type":
            mapping = {"P": "PAID", "R": "REVERSAL", "A": "ADJUSTMENT"}
            return mapping.get(self._row_status, "PAID")
        if field_type == "rx_number":
            return f"{rng.randint(100000000000, 999999999999)}"
        if field_type == "refill_number":
            return str(rng.randint(0, 11))
        if field_type == "line_number":
            return "1"
        if field_type == "days_supply":
            return str(rng.choice([30, 60, 90, 7, 14, 1]))
        if field_type == "quantity":
            return f"{rng.uniform(1.0, 360.0):.3f}"

        if field_type == "group_id":
            return bg["id"]
        if field_type == "group_name":
            return bg["name"]
        if field_type == "benefit_group":
            return bg["id"]
        if field_type == "segment_code":
            return self.maybe_null(f"SEG{rng.randint(1, 5):02d}")
        if field_type == "contract_number":
            return f"CTR{rng.randint(10000, 99999)}"
        if field_type == "carrier_id":
            return f"CAR{rng.randint(100, 999)}"
        if field_type == "customer_id":
            return f"CUS{rng.randint(10000, 99999)}"
        if field_type == "source_system":
            return self.pbm_key.upper()

        if field_type == "date_field":
            base = date(2024, 1, 1)
            d = base + timedelta(days=rng.randint(0, 730))
            return d.strftime("%Y-%m-%d")

        if field_type == "money":
            upper = field_name.upper()
            if "INGREDIENT" in upper or "BILLED" in upper:
                return f"{rng.uniform(5.0, 5000.0):.2f}"
            if "DISPENSING" in upper or "FEE" in upper:
                return f"{rng.uniform(0.0, 15.0):.2f}"
            if "COPAY" in upper or "COINSURANCE" in upper:
                return f"{rng.uniform(0.0, 75.0):.2f}"
            if "DEDUCTIBLE" in upper:
                return f"{rng.uniform(0.0, 500.0):.2f}"
            if "OOP" in upper:
                return f"{rng.uniform(0.0, 2000.0):.2f}"
            if "TAX" in upper:
                return f"{rng.uniform(0.0, 5.0):.2f}"
            if "PLAN" in upper or "PAID" in upper:
                return f"{rng.uniform(5.0, 4500.0):.2f}"
            return f"{rng.uniform(0.0, 500.0):.2f}"

        if field_type == "indicator":
            return rng.choice(["Y", "N"])
        if field_type == "zip":
            return f"{rng.randint(10000, 99999)}"
        if field_type == "phone":
            return f"{rng.randint(200, 999)}{rng.randint(2000000, 9999999)}"
        if field_type == "state":
            return rng.choice(STATES)
        if field_type == "city":
            return rng.choice(CITIES)
        if field_type == "address":
            return f"{rng.randint(1, 9999)} {rng.choice(STREETS)}"
        if field_type == "address2":
            return self.maybe_null(f"STE {rng.randint(100, 999)}", 0.70)
        if field_type == "code":
            return f"{rng.randint(0, 99):02d}"
        if field_type == "numeric_id":
            return str(rng.randint(100000, 9999999))
        if field_type == "name_field":
            return rng.choice(FIRST_NAMES + LAST_NAMES)
        if field_type == "text":
            return self.maybe_null("SAMPLE TEXT", 0.50)
        if field_type == "timestamp":
            base = date(2024, 1, 1)
            d = base + timedelta(days=rng.randint(0, 730))
            return f"{d.strftime('%Y-%m-%d')} {rng.randint(0,23):02d}:{rng.randint(0,59):02d}:{rng.randint(0,59):02d}"
        if field_type == "count":
            return str(rng.randint(0, 100))
        if field_type == "percent":
            return f"{rng.uniform(0, 100):.2f}"

        return f"VAL{rng.randint(1000, 9999)}"


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def read_file_definitions(excel_path):
    wb = openpyxl.load_workbook(str(excel_path), data_only=True)
    ws = wb["File column names"]

    files = {}
    for col_idx in range(2, 14):
        short_name = ws.cell(row=2, column=col_idx).value
        if not short_name:
            continue
        files[col_idx] = {
            "short_name": short_name,
            "fields": [],
        }

    for row_idx in range(3, 880):
        field_name = ws.cell(row=row_idx, column=1).value
        if not field_name:
            continue
        for col_idx in files:
            val = ws.cell(row=row_idx, column=col_idx).value
            if val and str(val).strip().upper() == "X":
                files[col_idx]["fields"].append(field_name)

    wb.close()
    return files


def generate_file(file_def, output_dir, rng, members, ndcs, pharmacies,
                  prescribers, hierarchies, benefit_groups):
    short_name = file_def["short_name"]
    pbm_key = PBM_SHORT_NAMES.get(short_name, short_name.lower().replace(" ", "_"))
    fields = file_def["fields"]

    field_types = [(f, classify_field(f)) for f in fields]

    gen = ValueGenerator(
        rng, members, ndcs, pharmacies, prescribers,
        hierarchies, benefit_groups, pbm_key,
    )

    filename = f"{pbm_key}.txt"
    filepath = output_dir / filename

    with open(filepath, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(DELIMITER.join(fields) + "\n")
        for _ in range(ROWS_PER_FILE):
            gen.new_row()
            row_values = []
            for field_name, field_type in field_types:
                row_values.append(str(gen.generate(field_name, field_type)))
            fh.write(DELIMITER.join(row_values) + "\n")

    return filename, len(fields)


def main():
    rng = random.Random(SEED)
    print(f"Reading file definitions from {EXCEL_PATH}")
    files = read_file_definitions(EXCEL_PATH)

    print("Building shared entity pools...")
    members, ndcs, pharmacies, prescribers, hierarchies, benefit_groups = (
        build_entity_pools(rng)
    )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    print(f"\nGenerating {ROWS_PER_FILE} rows per file into {OUTPUT_DIR}/\n")
    for col_idx in sorted(files.keys()):
        file_def = files[col_idx]
        filename, num_fields = generate_file(
            file_def, OUTPUT_DIR, rng, members, ndcs, pharmacies,
            prescribers, hierarchies, benefit_groups,
        )
        print(f"  {filename}: {num_fields} columns x {ROWS_PER_FILE} rows")

    print(f"\nDone. {len(files)} files written to {OUTPUT_DIR}/")


if __name__ == "__main__":
    main()
